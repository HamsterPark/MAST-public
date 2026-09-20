"""``/api/documents`` —— 文档浏览 / 编辑 / 版本 / 差异 / 认领 / 导出。

设计文档：``docs/v2/design/document_and_library_management.md`` §3.10

这个文件此前钉的是**文件名即身份**的世界（``draft:<title>_v003``、全局
``data/drafts/``）。那个世界有三个具体的谎：版本族由 LLM 自由填的 title 决定
（换个措辞就分叉、撞名就误合并）、拿着文件反查不出属于哪个实验、并发保存会
静默覆盖（``next_version_path`` 是 glob-then-write）。现在身份是 ``doc_id``、
落点是实验文件夹里的一个目录、版本分配只经 store。

所以这里断言的是**用户和智能体真正在意的性质**：

* 列表一行 = 一个文档（不是一个版本），带得出它属于哪个实验；
* 编辑 = 新版本，且**下一个打开该文档的智能体读到的就是它**（的
  全部意义 —— 旧编辑器把修改 POST 进一个没人读的内存 dict）；
* 别人抢先存了一版时，我的编辑拿到 409 而不是把他那一版盖掉；
* diff 能比**任意两版**（旧端点只能比相邻两版，「v1 到 v7 改了什么」问不出来）；
* HTML 导出的图是内嵌的（交付件永不断图），且留档不覆盖上一份。
"""
from __future__ import annotations

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from mast.api.context import AppContext
from mast.api.routes.documents import router
from mast.documents import KINDS, reset_caches, store
from mast.logging.storage import ExperimentStorage

# 1×1 透明 PNG —— 验证 HTML 导出真的把图 base64 嵌进去了。
_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
    "890000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082")


@pytest.fixture()
def env(tmp_path, monkeypatch):
    """一个空实验根 + 一条 active_scope 指针。

    ``reset_caches()`` 是必须的：store 的 ``doc_id → 目录`` 缓存是**进程内模块级**
    的，上一个测试的 tmp_path 会留在里面，不清就会在新的空目录里"找到"旧文档。
    """
    monkeypatch.setenv("MAST_EXPERIMENT_ROOT", str(tmp_path / "experiments"))
    monkeypatch.setenv("MAST_EXPERIMENT_DB", str(tmp_path / "db" / "mast_experiments.db"))
    st = ExperimentStorage(tmp_path / "db" / "mast_experiments.db")
    eid = st.create_experiment("Au111 形貌与 STS", "测 Au(111) 台阶")
    other = st.create_experiment("NbSe2 CDW", "另一个实验")
    st.set_active_scope(eid, None, updated_by="test")
    # 作用域解析优先读进程内 live ExperimentLog；单元测试里没有，但别的测试可能
    # 留下一个单例 —— 显式按掉，让作用域只来自 active_scope 表。
    import mast.logging.experiment_log as el
    monkeypatch.setattr(el, "get_active_log", lambda: None)
    reset_caches()
    yield {"storage": st, "eid": eid, "other": other, "tmp": tmp_path}
    reset_caches()


@pytest.fixture()
def client(env) -> TestClient:
    app = FastAPI()
    ctx = AppContext()
    ctx.wire(experiment_storage=env["storage"])
    app.state.ctx = ctx
    app.include_router(router, prefix="/api")
    return TestClient(app)


def _save(text: str, **kw):
    res = store().save(text=text, **kw)
    assert res.ok, res.error
    return res


# ════════════════════════════════════════════════════════════════════════
# 列表
# ════════════════════════════════════════════════════════════════════════

def test_empty_is_not_degraded(client):
    """还没有任何文档 —— 这是**空**，不是错误。"""
    b = client.get("/api/documents").json()
    assert b["degraded"] is False
    assert b["total"] == 0 and b["documents"] == []
    assert b["documents_root"], "空列表时至少要告诉用户文档会存到哪里"


def test_kinds_enumeration_is_the_single_source(client):
    """kind 清单**由服务端发下来**，前端不自己维护对照表。

    这是刻意不用 pydantic ``Literal`` 的原因：多一份白名单就是本项目反复踩的
    「一处定义、多处白名单」，而 ``Literal`` 漏一个 kind 会让一份磁盘上完好的
    文档在响应校验时炸掉、在 UI 里彻底消失。
    """
    b = client.get("/api/documents").json()
    assert [k["value"] for k in b["kinds"]] == list(KINDS)
    assert all(k["label"] for k in b["kinds"]), "每个 kind 都要有中文标签给徽章用"


def test_one_row_per_document_not_per_version(client, env):
    """两版是一个文档的历史，不是两份文档 —— 旧实现把每个 ``_vNNN.md`` 当一行。"""
    r = _save("# T\nv1", kind="experiment_report", title="Au111 报告")
    _save("# T\nv2", doc_id=r.doc_id)

    b = client.get("/api/documents").json()
    assert b["total"] == 1
    row = b["documents"][0]
    assert row["version"] == 2 and row["versions_count"] == 2
    assert row["doc_id"] == r.doc_id


def test_a_row_says_which_experiment_it_belongs_to(client, env):
    """「拿着 ``Au111_report_v003.md`` 反查不出它属于哪个实验」—— 修的就是这个。"""
    _save("# T\nbody", kind="experiment_report", title="Au111 报告")
    row = client.get("/api/documents").json()["documents"][0]
    assert row["experiment_id"] == env["eid"]
    assert row["experiment_name"] == "Au111 形貌与 STS"
    assert row["relation"] == "primary" and row["root_kind"] == "experiment"
    assert row["kind_label"] == "实验报告"
    assert row["path"].endswith("v001.md"), "path 要指向最新版本文件（用户复制它）"


def test_review_verdict_and_target_are_explicit(client, env):
    """评审的目标用 ``target_doc_id`` 明确记录，不再靠 draft_name 前缀猜。"""
    draft = _save("# 报告\nbody", kind="experiment_report", title="Au111 报告")
    _save("<!-- verdict: REVISE -->\n## 结论\n1. 补 STS 条件", kind="review",
          title="Au111 评审", target_doc_id=draft.doc_id, target_version=1)

    rev = next(d for d in client.get("/api/documents").json()["documents"]
               if d["kind"] == "review")
    assert rev["verdict"] == "REVISE"
    assert rev["target_doc_id"] == draft.doc_id and rev["target_version"] == 1


@pytest.mark.parametrize("params,expect", [
    ({"kind": "review"}, 1),
    ({"kind": "experiment_report"}, 1),
    ({"kind": "paper_draft"}, 0),
])
def test_kind_filter(client, env, params, expect):
    _save("# T\nbody", kind="experiment_report", title="报告")
    _save("<!-- verdict: ACCEPT -->\nok", kind="review", title="评审")
    assert client.get("/api/documents", params=params).json()["total"] == expect


def test_related_experiment_sees_the_document(client, env):
    """「关联实验的详情页也能看到它」—— 既定的归属模型（主归属 + 可关联）。"""
    r = _save("# T\nbody", kind="experiment_report", title="报告")
    client.patch(f"/api/documents/{r.doc_id}",
                 json={"related_experiment_ids": [env["other"]]})

    b = client.get("/api/documents", params={"experiment_id": env["other"]}).json()
    assert b["total"] == 1 and b["documents"][0]["relation"] == "related"

    only_primary = client.get("/api/documents", params={
        "experiment_id": env["other"], "include_related": False}).json()
    assert only_primary["total"] == 0


def test_unfiled_documents_are_listed_and_flagged(client, env):
    """无活跃实验时保存 → ``_unfiled``。**内容一个字都不丢，等着被认领。**

    这不是记账损失而是内容损失的问题：拒存的话 LLM 写好的整篇只活在对话流里。
    """
    env["storage"].set_active_scope(None, None, updated_by="test")
    reset_caches()
    r = _save("# 无主\nbody", kind="experiment_report", title="无主的")
    assert r.root_kind == "unfiled"

    row = next(d for d in client.get("/api/documents").json()["documents"]
               if d["doc_id"] == r.doc_id)
    assert row["root_kind"] == "unfiled" and row["relation"] == "unfiled"
    assert row["experiment_id"] is None

    hidden = client.get("/api/documents", params={"include_unfiled": False}).json()
    assert all(d["doc_id"] != r.doc_id for d in hidden["documents"])


# ════════════════════════════════════════════════════════════════════════
# 详情 / 版本
# ════════════════════════════════════════════════════════════════════════

def test_detail_serves_latest_and_any_version(client, env):
    r = _save("# T\nFIRST", kind="experiment_report", title="报告")
    _save("# T\nSECOND", doc_id=r.doc_id)

    d = client.get(f"/api/documents/{r.doc_id}").json()
    assert "SECOND" in d["content"] and d["body"] == d["content"]
    assert len(d["versions"]) == 2 and d["versions"][0]["sha256"]

    old = client.get(f"/api/documents/{r.doc_id}", params={"version": 1}).json()
    assert "FIRST" in old["content"] and "SECOND" not in old["content"]
    assert old["version_requested"] == 1


def test_current_resolves_to_the_latest_document(client, env):
    """``doc_id=current`` —— 对话里记下的「当前那一份」也能直接打开。"""
    _save("# 旧\nbody", kind="experiment_report", title="旧报告")
    newer = _save("# 新\nbody", kind="experiment_report", title="新报告")
    assert client.get("/api/documents/current").json()["doc_id"] == newer.doc_id


def test_unknown_doc_id_is_a_structured_404(client):
    r = client.get("/api/documents/01JZZZNOPE")
    assert r.status_code == 404
    assert r.json()["detail"], "404 必须说清是找不到，而不是空响应"
    assert r.json()["content"] == ""


def test_versions_include_an_orphan_version_file(client, env):
    """**版本集合 = 目录扫描 ∪ versions.jsonl**（读侧自愈，设计陷阱 ⑬）。

    claim 成功但 jsonl 追加前崩了的版本仍然是一个真实存在的、有内容的版本；
    只信 jsonl 就等于在崩溃后静默丢掉用户刚写的东西。
    """
    r = _save("# T\nv1", kind="experiment_report", title="报告")
    entry = store().get(r.doc_id)
    (entry.dir / "v002.md").write_text("# T\n崩在登记前的一版", encoding="utf-8")
    (entry.dir / "v003.md").write_text("", encoding="utf-8")  # 空 claim：不算一版

    v = client.get(f"/api/documents/{r.doc_id}/versions").json()
    assert [x["version"] for x in v["versions"]] == [1, 2]
    assert v["latest_version"] == 2


def test_versions_carry_provenance(client, env):
    r = _save("# T\nbody", kind="experiment_report", title="报告",
              created_by="agent:paper_writing", note="首版")
    v = client.get(f"/api/documents/{r.doc_id}/versions").json()
    row = v["versions"][0]
    assert row["created_by"] == "agent:paper_writing" and row["note"] == "首版"
    assert row["sha256"] and row["words"] > 0


# ════════════════════════════════════════════════════════════════════════
# diff —— 任意两版
# ════════════════════════════════════════════════════════════════════════

class TestDiff:
    def test_any_two_versions_not_just_adjacent(self, client, env):
        """旧端点只能比 ``versions[-2]`` vs ``versions[-1]``，「v1 到 v3 一共改了
        什么」根本问不出来 —— 而那正是用户想问的。"""
        r = _save("# T\nline one\nline two", kind="experiment_report", title="报告")
        _save("# T\nline one\nline two CHANGED", doc_id=r.doc_id)
        _save("# T\nline one\nline two CHANGED\nline three", doc_id=r.doc_id)

        d = client.get(f"/api/documents/{r.doc_id}/diff",
                       params={"from_v": 1, "to_v": 3}).json()
        assert d["from_version"] == 1 and d["to_version"] == 3 and d["changed"]
        assert "-line two" in d["diff"] and "+line two CHANGED" in d["diff"]
        assert "+line three" in d["diff"]
        assert d["versions_available"] == [1, 2, 3]

    def test_default_is_the_latest_two(self, client, env):
        r = _save("# T\na", kind="experiment_report", title="报告")
        _save("# T\nb", doc_id=r.doc_id)
        _save("# T\nc", doc_id=r.doc_id)
        d = client.get(f"/api/documents/{r.doc_id}/diff").json()
        assert (d["from_version"], d["to_version"]) == (2, 3)

    def test_a_lone_version_is_not_reported_as_all_new(self, client, env):
        """旧 diff 拿空串当 original，于是每次都把整篇报成新增。"""
        r = _save("# T\nonly", kind="experiment_report", title="报告")
        d = client.get(f"/api/documents/{r.doc_id}/diff").json()
        assert d["changed"] is False and d["diff"] == ""
        assert d["detail"], "没有可比的版本时要说明原因，而不是给个空 diff"

    @pytest.mark.parametrize("params", [{"to_v": 99}, {"from_v": 99, "to_v": 1}])
    def test_nonexistent_version_is_a_404_naming_what_exists(self, client, env, params):
        r = _save("# T\nonly", kind="experiment_report", title="报告")
        resp = client.get(f"/api/documents/{r.doc_id}/diff", params=params)
        assert resp.status_code == 404
        assert "现有" in resp.json()["detail"]


# ════════════════════════════════════════════════════════════════════════
# PUT —— 用户编辑（的核心）
# ════════════════════════════════════════════════════════════════════════

class TestOperatorEdit:
    def test_edit_is_a_new_version_and_reaches_the_agent(self, client, env):
        """旧对象编辑器把修改 POST 进内存 ``artifact_edits``，docstring 声称
        「orchestrator 下一个 super-step 会读」—— **从来没有任何 agent 读过它**。
        写进文件才闭环：评审 agent 的 ``load_draft`` 打开的就是这个文件。"""
        from mast.agents.paper_review.tools import load_draft

        r = _save("# T\nAGENT ORIGINAL", kind="experiment_report", title="Au111 报告",
                  created_by="agent:paper_writing")
        resp = client.put(f"/api/documents/{r.doc_id}",
                          json={"content": "# T\nOPERATOR CORRECTED THIS"})
        b = resp.json()
        assert resp.status_code == 200 and b["ok"]
        assert b["version"] == 2, "编辑是新版本，不是覆盖"

        assert "OPERATOR CORRECTED THIS" in load_draft.invoke({"draft_id": r.doc_id})
        # 智能体自己那一版还在盘上
        v = client.get(f"/api/documents/{r.doc_id}/versions").json()
        assert v["total"] == 2 and v["versions"][-1]["created_by"] == "operator"

    def test_note_lands_in_the_version_log(self, client, env):
        r = _save("# T\nbody", kind="experiment_report", title="报告")
        client.put(f"/api/documents/{r.doc_id}",
                   json={"content": "# T\n改过的", "note": "修正单位错误"})
        v = client.get(f"/api/documents/{r.doc_id}/versions").json()
        assert v["versions"][-1]["note"] == "修正单位错误"

    def test_stale_base_version_is_a_409_and_writes_nothing(self, client, env):
        """智能体在我编辑期间又存了一版 —— 我的文本不能把它盖掉。

        版本本身不会丢（新版本永不覆盖旧版本），但「基于过期内容的编辑成为最新
        版本」会让智能体下次读到一份倒退的正文，那是实打实的内容损失。
        """
        r = _save("# T\nv1", kind="experiment_report", title="报告")
        _save("# T\nv2 智能体刚存的", doc_id=r.doc_id)  # 我看到的是 v1

        resp = client.put(f"/api/documents/{r.doc_id}",
                          json={"content": "# T\n基于 v1 的编辑", "base_version": 1})
        assert resp.status_code == 409
        b = resp.json()
        assert b["conflict"] is True and b["latest_version"] == 2
        assert "已有更新的版本" in b["detail"]
        assert client.get(f"/api/documents/{r.doc_id}/versions").json()["total"] == 2

    def test_matching_base_version_goes_through(self, client, env):
        r = _save("# T\nv1", kind="experiment_report", title="报告")
        resp = client.put(f"/api/documents/{r.doc_id}",
                          json={"content": "# T\n改过的", "base_version": 1})
        assert resp.status_code == 200 and resp.json()["version"] == 2

    def test_empty_content_refused(self, client, env):
        r = _save("# T\nbody", kind="experiment_report", title="报告")
        resp = client.put(f"/api/documents/{r.doc_id}", json={"content": "   "})
        assert resp.status_code == 400 and resp.json()["ok"] is False
        assert client.get(f"/api/documents/{r.doc_id}/versions").json()["total"] == 1

    def test_legacy_body_field_still_writes(self, client, env):
        """旧前端发的是 ``{body: …}``。切换期间两个字段名都得能写。"""
        r = _save("# T\nbody", kind="experiment_report", title="报告")
        resp = client.put(f"/api/documents/{r.doc_id}", json={"body": "# T\n旧字段写入"})
        assert resp.status_code == 200 and resp.json()["version"] == 2

    def test_unknown_doc_id_is_404_not_a_new_document(self, client, env):
        """写一个不存在的 doc_id **绝不能**悄悄新建一份 —— 那样用户的修改就
        落在一份没人会打开的孤立文档里。"""
        resp = client.put("/api/documents/01JZZZNOPE", json={"content": "x"})
        assert resp.status_code == 404
        assert client.get("/api/documents").json()["total"] == 0


# ════════════════════════════════════════════════════════════════════════
# PATCH —— 可变头部
# ════════════════════════════════════════════════════════════════════════

class TestPatch:
    def test_rename_does_not_create_a_version_or_change_identity(self, client, env):
        """标题只是显示名。改名不发版本、不动目录、不换 doc_id —— 「title 当族键」
        正是旧实现分叉与误合并的根因（设计陷阱 ②）。"""
        r = _save("# T\nbody", kind="experiment_report", title="旧标题")
        before = store().get(r.doc_id).dir.name

        resp = client.patch(f"/api/documents/{r.doc_id}", json={"title": "新标题"})
        assert resp.status_code == 200 and resp.json()["title"] == "新标题"
        assert client.get(f"/api/documents/{r.doc_id}/versions").json()["total"] == 1
        assert store().get(r.doc_id).dir.name == before, "目录名创建时冻结"
        assert store().get(r.doc_id).meta.title_history, "旧标题要留痕"

    def test_kind_correction(self, client, env):
        """kind 标错由用户更正；``doc.json`` 才是权威，目录名前缀不追改。"""
        r = _save("# T\nbody", kind="experiment_report", title="其实是投稿手稿")
        resp = client.patch(f"/api/documents/{r.doc_id}", json={"kind": "paper_draft"})
        assert resp.json()["kind"] == "paper_draft"
        assert resp.json()["kind_label"] == "论文草稿"

    def test_unknown_doc_id_is_404(self, client):
        assert client.patch("/api/documents/01JZZZNOPE",
                            json={"title": "x"}).status_code == 404


# ════════════════════════════════════════════════════════════════════════
# claim —— 认领 / 改主归属（物理搬家）
# ════════════════════════════════════════════════════════════════════════

class TestClaim:
    def test_unfiled_document_can_be_claimed(self, client, env):
        env["storage"].set_active_scope(None, None, updated_by="test")
        reset_caches()
        r = _save("# 无主\nbody", kind="experiment_report", title="无主的")
        assert r.root_kind == "unfiled"

        resp = client.post(f"/api/documents/{r.doc_id}/claim",
                           json={"experiment_id": env["eid"]})
        b = resp.json()
        assert resp.status_code == 200 and b["ok"] and b["root_kind"] == "experiment"
        assert b["experiment_id"] == env["eid"]

        entry = store().get(r.doc_id)
        assert entry.dir.is_dir() and "reports" in entry.dir.parts
        assert entry.read_text() and "body" in entry.read_text(), "搬家不能丢内容"

        listed = client.get("/api/documents", params={
            "experiment_id": env["eid"], "include_related": False}).json()
        assert [d["doc_id"] for d in listed["documents"]] == [r.doc_id]

    def test_claiming_into_a_nonexistent_experiment_is_a_400(self, client, env):
        r = _save("# T\nbody", kind="experiment_report", title="报告")
        resp = client.post(f"/api/documents/{r.doc_id}/claim",
                           json={"experiment_id": "no-such-experiment"})
        assert resp.status_code == 400 and resp.json()["ok"] is False
        assert resp.json()["detail"]

    def test_unknown_doc_id_is_404(self, client, env):
        assert client.post("/api/documents/01JZZZNOPE/claim",
                           json={"experiment_id": env["eid"]}).status_code == 404


# ════════════════════════════════════════════════════════════════════════
# export
# ════════════════════════════════════════════════════════════════════════

class TestExport:
    def test_markdown_download_is_named_after_the_document(self, client, env):
        """旧 ArtifactEditor 硬编码 ``${id}.txt``，下载一堆认不出的文件。"""
        r = _save("# T\nMARKER", kind="experiment_report", title="Au111 形貌与 STS")
        resp = client.get(f"/api/documents/{r.doc_id}/export", params={"format": "md"})
        assert resp.status_code == 200 and "MARKER" in resp.text
        cd = resp.headers["content-disposition"]
        assert "attachment" in cd and "filename*=UTF-8''" in cd and "_v001" in cd

    def test_html_inlines_the_figures(self, client, env):
        """自包含 HTML 是唯一的交付格式：md 里的 ``../_assets/x.png`` 一离开实验
        文件夹就断，而 base64 内嵌的图**永不断**。

        ``base_dir`` 必须是版本文件所在目录 —— 传实验目录或 cwd 会让每张图都
        missing，而那会静默交付一份没有数据的报告。
        """
        from mast.documents import paths as dpaths

        home, root_kind = dpaths.doc_home("experiment_report", env["eid"], create=True,
                                          st=env["storage"])
        assets = dpaths.assets_dir_for(home, root_kind, create=True)
        (assets / "topo.png").write_bytes(_PNG)
        r = _save("# 报告\n![形貌](../_assets/topo.png)\n", kind="experiment_report",
                  title="Au111 形貌")

        resp = client.get(f"/api/documents/{r.doc_id}/export", params={"format": "html"})
        assert resp.status_code == 200 and "text/html" in resp.headers["content-type"]
        assert "data:image/png;base64," in resp.text
        assert resp.headers["X-Images-Inlined"] == "1"
        assert resp.headers["X-Images-Missing"] == "0"

    def test_html_export_is_archived_without_overwriting(self, client, env):
        """留档带时间戳、多份共存。旧实现固定文件名直接覆盖，上一份交付件就没了。"""
        from urllib.parse import unquote

        r = _save("# T\nbody", kind="experiment_report", title="报告")
        first = client.get(f"/api/documents/{r.doc_id}/export", params={"format": "html"})
        p1 = unquote(first.headers["X-Export-Path"])
        assert p1 and "exports" in p1

        _save("# T\nbody v2", doc_id=r.doc_id)
        second = client.get(f"/api/documents/{r.doc_id}/export", params={"format": "html"})
        p2 = unquote(second.headers["X-Export-Path"])
        from pathlib import Path
        assert p2 != p1 and Path(p1).is_file() and Path(p2).is_file()
        assert "_v001_" in Path(p1).name and "_v002_" in Path(p2).name

    def test_docx_is_a_real_word_file_with_builtin_styles(self, client, env):
        """docx 覆盖的是 HTML 覆盖不了的那半:发给别人**改**。

        所以断言的不是「导出成功」，而是收件人打开后能用 Word 的那套东西 ——
        结构挂内置样式（改一次 ``Heading 1`` 全文重排、导航窗格、自动目录、修订
        批注都依赖它），图真嵌进包里。渲染器本身的样式映射在
        ``tests/v2/unit/documents/test_report_docx.py`` 里逐条钉。
        """
        import io

        import docx

        from mast.documents import paths as dpaths

        home, root_kind = dpaths.doc_home("experiment_report", env["eid"], create=True,
                                          st=env["storage"])
        assets = dpaths.assets_dir_for(home, root_kind, create=True)
        (assets / "topo.png").write_bytes(_PNG)
        r = _save("# 阶段总结\n\n偏压 **-0.5 V**。\n\n![形貌](../_assets/topo.png)\n",
                  kind="experiment_report", title="Au111 形貌")

        resp = client.get(f"/api/documents/{r.doc_id}/export", params={"format": "docx"})
        assert resp.status_code == 200
        assert "wordprocessingml.document" in resp.headers["content-type"]
        assert resp.content[:2] == b"PK", "docx 是 zip 容器，头两字节必须是 PK"
        assert resp.headers["X-Images-Inlined"] == "1"
        assert resp.headers["X-Images-Missing"] == "0"

        d = docx.Document(io.BytesIO(resp.content))
        assert [p.style.name for p in d.paragraphs if p.text.strip()][:3] == [
            "Title", "Subtitle", "Heading 1"]
        # 副标题带版本号：一份 docx 一旦离开系统就没有 sidecar，文件里必须自己
        # 说清是哪一版，否则收件人手上两份改稿分不出先后。
        assert "v001" in d.paragraphs[1].text
        assert len(d.inline_shapes) == 1

    def test_docx_export_is_archived_with_a_docx_suffix(self, client, env):
        """留档扩展名跟着格式走 —— ``exports/`` 里 ``.html`` 与 ``.docx`` 并存。"""
        from pathlib import Path
        from urllib.parse import unquote

        r = _save("# T\nbody", kind="experiment_report", title="报告")
        html_p = Path(unquote(client.get(
            f"/api/documents/{r.doc_id}/export",
            params={"format": "html"}).headers["X-Export-Path"]))
        docx_p = Path(unquote(client.get(
            f"/api/documents/{r.doc_id}/export",
            params={"format": "docx"}).headers["X-Export-Path"]))
        assert html_p.suffix == ".html" and docx_p.suffix == ".docx"
        assert html_p.is_file() and docx_p.is_file()

    def test_word_is_accepted_as_an_alias(self, client, env):
        """要求的是「word」，不是「docx」。"""
        r = _save("# T\nbody", kind="experiment_report", title="报告")
        resp = client.get(f"/api/documents/{r.doc_id}/export", params={"format": "word"})
        assert resp.status_code == 200 and resp.content[:2] == b"PK"

    def test_missing_python_docx_says_so_and_names_the_fix(self, client, env,
                                                           monkeypatch):
        """python-docx 缺失（未装 / 打包漏了 ``docx/templates/`` 数据）时不许含糊。

        这是一个**真会发生**的场景：冻结包若没 ``collect_all("docx")``，
        ``Document()`` 就找不到默认模板。那时错误必须说清「没装」并给出装的命令，
        而不是长成一句「渲染失败」让人去查 markdown。md / html 不受影响也要说，
        否则用户以为导出整个坏了。
        """
        import builtins

        real_import = builtins.__import__

        def _no_docx(name, *a, **kw):
            if name == "docx" or name.startswith("docx."):
                raise ImportError("No module named 'docx'")
            return real_import(name, *a, **kw)

        r = _save("# T\nbody", kind="experiment_report", title="报告")
        monkeypatch.delitem(__import__("sys").modules, "docx", raising=False)
        monkeypatch.delitem(__import__("sys").modules,
                            "mast.agents._shared.report_docx", raising=False)
        monkeypatch.setattr(builtins, "__import__", _no_docx)
        resp = client.get(f"/api/documents/{r.doc_id}/export", params={"format": "docx"})

        assert resp.status_code == 400
        body = resp.json()
        assert body["degraded"] is True
        assert "python-docx" in body["detail"] and "未安装" in body["detail"]
        assert "md" in body["detail"] and "html" in body["detail"]

    def test_unsupported_format_is_a_400_naming_what_works(self, client, env):
        r = _save("# T\nbody", kind="experiment_report", title="报告")
        resp = client.get(f"/api/documents/{r.doc_id}/export", params={"format": "pdf"})
        assert resp.status_code == 400
        detail = resp.json()["detail"]
        assert "md" in detail and "html" in detail and "docx" in detail

    def test_unknown_doc_id_is_404(self, client):
        assert client.get("/api/documents/01JZZZNOPE/export").status_code == 404


# ════════════════════════════════════════════════════════════════════════
# 没有终态
# ════════════════════════════════════════════════════════════════════════

def test_there_is_no_finalize_or_archive_endpoint(client):
    """文档没有终态（INCREMENTAL-ONLY，与 ``experiment.json`` 同一条铁律）。

    十年后重新打开一个实验、再写一版报告必须是合法操作，所以这里**刻意**没有
    finalize / archive / close 之类的端点，也没有 ``status`` 字段。
    """
    paths = {r.path for r in client.app.routes if hasattr(r, "path")}
    assert not [p for p in paths
                if any(w in p for w in ("finalize", "archive", "close"))]
    from mast.api.schemas_documents import DocumentSummary
    fields = set(DocumentSummary.model_fields)
    assert not fields & {"status", "finalized_at", "archived"}
