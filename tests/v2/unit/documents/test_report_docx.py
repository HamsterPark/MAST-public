"""Word 导出:全部挂 **Word 内置样式**,图嵌不进去要如实说。

为什么有这条路(设计 §6 的更正):HTML 自包含导出覆盖了「发给别人**看**」,但覆盖不了
「让别人**改**」—— 导师/合作者要用 Word 的修订与批注,期刊也不收 HTML 投稿。

为什么钉「内置样式」而不是钉外观:内置样式是**可编辑性**的载体。收件人改一次
``Heading 1`` 就能重排全文;修订/审阅、导航窗格、自动目录、交叉引用全都依赖它;
套期刊模板只是一次样式替换。一份用「放大加粗的正文」冒充标题的 docx,在 Word 里是
一份没有结构的文档 —— 打开看着像,用起来全不对。
"""

from __future__ import annotations

import struct
import zlib
from pathlib import Path

import pytest

from mast.agents._shared.report_docx import render_docx


def _png(w: int = 8, h: int = 6) -> bytes:
    """一张**合法**的 PNG(逐 chunk CRC 正确)。

    这里自己造而不借现成常量,是因为仓库里两串「minimal 1x1 PNG」并不都合法:
    ``tests/v2/unit/api/test_documents.py:32`` 的 ``_PNG`` 三个 chunk CRC 全对、
    python-docx 收;``tests/v2/agents/test_paper_xd_honesty_fixes.py:138`` 的
    ``_PNG_A`` 的 IDAT 长度字段(13)小于实际载荷(16),解析器按错位置读下一个
    chunk 头 → python-docx 抛 ``UnicodeDecodeError``。那个常量只用来比 sha256
    (测同名不同字节不静默覆盖),从不真解码,所以一直没暴露。

    结论不是「那个测试有错」,而是:**PNG 合法性不能靠肉眼**。浏览器宽容,
    base64 内嵌进 HTML 照样"成功",坏图要到读者打开时才发现。
    """
    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return (struct.pack(">I", len(data)) + body
                + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF))

    raw = b"".join(b"\x00" + b"\x20\x60\xa0" * w for _ in range(h))
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b""))


@pytest.fixture()
def layout(tmp_path) -> tuple[Path, Path]:
    """复刻真实布局:图池在 ``<exp>/reports/_assets/``，文档在 ``<exp>/reports/<doc>/``。

    布局搭错(把 ``_assets`` 放到 ``reports/`` 外面)会让每张图都 missing —— 那时该
    先怀疑测量设置，而不是代码。
    """
    assets = tmp_path / "reports" / "_assets"
    assets.mkdir(parents=True)
    doc_dir = tmp_path / "reports" / "doc"
    doc_dir.mkdir(parents=True)
    return doc_dir, assets


def _open(blob: bytes, tmp_path: Path):
    import docx
    p = tmp_path / "out.docx"
    p.write_bytes(blob)
    return docx.Document(str(p))


def _styles(doc) -> list[str]:
    return [p.style.name for p in doc.paragraphs if p.text.strip()]


def test_structure_uses_word_builtin_paragraph_styles(layout, tmp_path):
    """★ 标题/列表/引用/代码全部挂内置段落样式，不是「放大加粗的正文」。"""
    doc_dir, _ = layout
    md = (
        "# 一级标题\n\n正文一句。\n\n## 二级标题\n\n"
        "1. 有序一\n2. 有序二\n\n- 无序一\n- 无序二\n\n"
        "> 一句引用\n\n```\ncode_line_1()\ncode_line_2()\n```\n"
    )
    blob, _emb, _miss = render_docx(md, base_dir=doc_dir, title="标题", subtitle="副标题")
    d = _open(blob, tmp_path)
    assert _styles(d) == [
        "Title", "Subtitle", "Heading 1", "Normal", "Heading 2",
        "List Number", "List Number", "List Bullet", "List Bullet",
        "Quote", "No Spacing", "No Spacing",
    ]


def test_inline_markup_uses_builtin_character_styles(layout, tmp_path):
    """粗/斜/代码用 ``Strong``/``Emphasis``/``Macro Text Char``，而不是 ``run.bold``。

    区别在于可编辑性：字符样式能被一次性重定义，``run.bold`` 是死格式。
    """
    doc_dir, _ = layout
    blob, _e, _m = render_docx(
        "偏压 **-0.5 V**、探针 *W*、命令 `Scan_Action()`。",
        base_dir=doc_dir, title="")
    d = _open(blob, tmp_path)
    styled = [(r.style.name, r.text) for p in d.paragraphs for r in p.runs
              if r.style and r.style.name != "Default Paragraph Font"]
    assert styled == [("Strong", "-0.5 V"), ("Emphasis", "W"),
                      ("Macro Text Char", "Scan_Action()")]


def test_deeper_headings_clamp_to_heading_4(layout, tmp_path):
    doc_dir, _ = layout
    blob, _e, _m = render_docx("##### 五级\n\n###### 六级\n", base_dir=doc_dir, title="")
    assert _styles(_open(blob, tmp_path)) == ["Heading 4", "Heading 4"]


def test_table_uses_table_grid_with_strong_header(layout, tmp_path):
    """``Table Grid`` 是最朴素、最好改的内置表格样式 —— 这份文档的用途是被人改。"""
    doc_dir, _ = layout
    md = "| 参数 | 值 |\n|---|---|\n| 偏压 | -0.5 |\n| 电流 | 100 |\n"
    blob, _e, _m = render_docx(md, base_dir=doc_dir, title="")
    d = _open(blob, tmp_path)
    assert len(d.tables) == 1
    t = d.tables[0]
    assert t.style.name == "Table Grid"
    # 表头每个单元格恰好一个 run 且挂 Strong（不该有 cell.text="" 造出的空 run）
    header_runs = [r for c in t.rows[0].cells for r in c.paragraphs[0].runs]
    assert [r.style.name for r in header_runs] == ["Strong", "Strong"]
    assert [r.text for r in header_runs] == ["参数", "值"]
    assert [c.text for c in t.rows[1].cells] == ["偏压", "-0.5"]


def test_image_is_embedded_with_caption_style_and_width_capped(layout, tmp_path):
    """图嵌进文档、题注用 ``Caption``、宽度不超过正文可用宽。

    题注用 Word 的 ``Caption`` 样式不是为了好看：它能被「插入题注」体系接管，
    也就是收件人可以据此自动编号和做交叉引用。
    """
    doc_dir, assets = layout
    (assets / "topo.png").write_bytes(_png(64, 48))
    md = "![形貌图](../_assets/topo.png)\n\n*Caption:* 图 1：Au(111) 人字纹重构\n"
    blob, embedded, missing = render_docx(md, base_dir=doc_dir, title="")
    assert (embedded, missing) == (1, 0)
    d = _open(blob, tmp_path)
    assert _styles(d) == ["Caption"]
    assert d.paragraphs[-1].text == "图 1：Au(111) 人字纹重构"
    assert len([r for r in d.part.rels.values() if "image" in r.reltype]) == 1
    sec = d.sections[0]
    usable = sec.page_width - sec.left_margin - sec.right_margin
    assert d.inline_shapes[0].width <= usable, "图宽超过正文宽会溢出到页面外"


@pytest.mark.parametrize("name,payload,expect", [
    ("nope.png", None, "图片缺失"),
    ("curve.pdf", b"%PDF-1.4 x", "Word 不支持嵌入"),
    ("empty.png", b"", "图片为空"),
])
def test_unembeddable_images_are_reported_in_the_document(layout, tmp_path,
                                                          name, payload, expect):
    """★ 图嵌不进去必须**写进正文**并计入 missing。

    静默消失的图正是「报告看着完整、实则少了数据」的来源 —— 那比报错更危险。
    """
    doc_dir, assets = layout
    if payload is not None:
        (assets / name).write_bytes(payload)
    blob, embedded, missing = render_docx(
        f"![图](../_assets/{name})\n", base_dir=doc_dir, title="")
    assert (embedded, missing) == (0, 1)
    d = _open(blob, tmp_path)
    body = "\n".join(p.text for p in d.paragraphs)
    assert expect in body


def test_remote_images_are_not_fetched(layout, tmp_path):
    """远程图不联网抓：仪器机常离线，一次挂住的 HTTP 请求比一张缺图更糟。"""
    doc_dir, _ = layout
    blob, embedded, missing = render_docx(
        "![远程](https://example.com/a.png)\n", base_dir=doc_dir, title="")
    assert (embedded, missing) == (0, 1)
    assert "外部图片未嵌入" in "\n".join(p.text for p in _open(blob, tmp_path).paragraphs)


def test_links_become_real_clickable_hyperlinks(layout, tmp_path):
    """链接要真能点 —— 可点击是实质，颜色是外观。

    默认模板里**没有** ``Hyperlink`` 字符样式（已核对），所以 run 不挂样式；
    自己造一个同名样式去冒充内置的，才是错的做法。
    """
    doc_dir, _ = layout
    blob, _e, _m = render_docx(
        "参考 [Barth 1990](https://doi.org/10.1103/PhysRevB.42.9307)。",
        base_dir=doc_dir, title="")
    d = _open(blob, tmp_path)
    links = [r for r in d.part.rels.values() if "hyperlink" in r.reltype]
    assert len(links) == 1
    assert links[0].target_ref.startswith("https://doi.org/")
    assert "Barth 1990" in d.paragraphs[0].text


def test_never_sets_hard_formatting(layout, tmp_path):
    """★ 从不硬编码字号/字体/颜色 —— 否则收件人改样式改不动全文。

    这是用户「用 word 内置」那句话的可执行版本。
    """
    doc_dir, _ = layout
    md = "# 标题\n\n正文 **粗** 和 `代码`。\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
    blob, _e, _m = render_docx(md, base_dir=doc_dir, title="T", subtitle="S")
    d = _open(blob, tmp_path)
    for p in d.paragraphs:
        for r in p.runs:
            assert r.font.size is None, f"run {r.text!r} 硬编码了字号"
            assert r.font.name is None, f"run {r.text!r} 硬编码了字体"
            assert r.font.color.rgb is None, f"run {r.text!r} 硬编码了颜色"


def test_unrecognised_content_degrades_to_a_paragraph(layout, tmp_path):
    """认不出来的东西当段落输出 —— 降级样式，**绝不丢内容**。"""
    doc_dir, _ = layout
    weird = "<<< 这不是任何 markdown 结构 >>> 但它必须出现在文档里"
    blob, _e, _m = render_docx(weird + "\n", base_dir=doc_dir, title="")
    body = "\n".join(p.text for p in _open(blob, tmp_path).paragraphs)
    assert weird in body


def test_empty_markdown_still_produces_a_valid_document(layout, tmp_path):
    doc_dir, _ = layout
    blob, embedded, missing = render_docx("", base_dir=doc_dir, title="只有标题")
    assert (embedded, missing) == (0, 0)
    assert _styles(_open(blob, tmp_path)) == ["Title"]


def test_every_style_we_rely_on_really_is_builtin():
    """★ 我们用的样式必须真的在**默认模板**里，且类型对得上。

    这条不测我们的代码，测的是**前提**：如果 python-docx 换版本后默认模板变了
    （少了某个样式、或段落/字符类型变了），渲染器会静默回落到 ``Normal`` —— 导出
    的 docx 打开看着差不多，实际上没有结构，收件人改样式改不动全文。那种失败不会
    有异常，只会有一份"看着像"的文档。

    也顺手钉住两个**不存在**的样式：``Macro Text``（只有字符版 ``Macro Text Char``）
    和 ``Hyperlink``。哪天它们出现了，是可以改用的信号；而在它们不存在时自己造一个
    同名样式去冒充内置，就把「用 word 内置」这条约束偷偷废掉了。
    """
    import docx
    from docx.enum.style import WD_STYLE_TYPE

    styles = {s.name: s.type for s in docx.Document().styles}
    P, C = WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER
    for name, kind in [
        ("Title", P), ("Subtitle", P), ("Heading 1", P), ("Heading 2", P),
        ("Heading 3", P), ("Heading 4", P), ("Normal", P), ("No Spacing", P),
        ("List Bullet", P), ("List Number", P), ("Quote", P), ("Caption", P),
        ("Strong", C), ("Emphasis", C), ("Macro Text Char", C),
    ]:
        assert styles.get(name) == kind, (
            f"默认模板里没有 {name}（或类型不是 {kind}）—— 渲染器会静默回落到 "
            f"Normal，导出的文档将失去结构")
    assert "Table Grid" in styles

    assert "Macro Text" not in styles, (
        "默认模板现在有 Macro Text 段落样式了 —— 代码块可以从 No Spacing 改用它")
    assert "Hyperlink" not in styles, (
        "默认模板现在有 Hyperlink 字符样式了 —— 链接 run 可以挂上它")


def test_verdict_marker_survives(layout, tmp_path):
    """评审文档首行的 ``<!-- verdict: X -->`` 是机器可读标记，不能被吞掉。"""
    doc_dir, _ = layout
    blob, _e, _m = render_docx("<!-- verdict: REVISE -->\n\n1. 缺误差棒\n",
                               base_dir=doc_dir, title="")
    body = "\n".join(p.text for p in _open(blob, tmp_path).paragraphs)
    assert "verdict: REVISE" in body and "缺误差棒" in body
